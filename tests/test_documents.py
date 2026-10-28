import io,json,zipfile
from pathlib import Path
from uuid import uuid4
import pytest
from reviewlab.services.documents import Documents,extract,packed,digest
from reviewlab.services.document_transfer import export_bundle,preview_bundle,restore_bundle
from reviewlab.analysis.documents import compare,suggestions
from reviewlab.services.document_ai import prepare,analyze,reserve_budget
from reviewlab.repositories.library import Library


def pdf(text=False,encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject,NameObject,DictionaryObject
    w=PdfWriter();p=w.add_blank_page(600,800)
    if text:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        p[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 700 Td (Review trust study.) Tj ET');p[NameObject('/Contents')]=w._add_object(stream)
    if encrypted:w.encrypt('not-real-password')
    out=io.BytesIO();w.write(out);return out.getvalue()


def test_pdf_scan_encrypted_and_extract():
    assert extract(pdf(), 'pdf')[0]=='ocr_needed'
    assert extract(pdf(encrypted=True),'pdf')[0]=='encrypted'
    status,units=extract(pdf(True),'pdf');assert status=='ready' and units[0]['location']=='페이지 1' and 'trust' in units[0]['text']
    assert extract(b'%PDF-broken','pdf')[0]=='extraction_failed'


def test_docx_korean_paragraph_and_table():
    from docx import Document
    doc=Document();doc.add_paragraph('가상 연구: 리뷰 신뢰도');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='표본';table.cell(0,1).text='가상 10명'
    out=io.BytesIO();doc.save(out)
    status,units=extract(out.getvalue(),'docx')
    assert status=='ready' and units[0]['location']=='문단 1' and units[1]['location']=='표 1 행 1'
    assert '가상 10명' in units[1]['text']


@pytest.mark.parametrize('blob,ext',[(b'','txt'),(b'\xff','txt'),(b'a\0b','txt'),(b'notpdf','pdf'),(b'zip','docx'),(b'abc','exe')])
def test_bad_documents_rejected(blob,ext):
    with pytest.raises(ValueError):extract(blob,ext)


def seed(library,source):
    from reviewlab.services.research import Research
    sid=Research(library).save(source);service=Documents(library)
    row=service.upload('가상 연구 목적: 리뷰 신뢰도 탐색\n결과: 표본 10명에서 차이가 관찰됨.'.encode(),'가상.txt',sid)
    return service,row


def test_search_evidence_history_duplicate_and_missing_file(library,source):
    s,d=seed(library,source);scope=[d['id']];hits=s.search(scope,'리뷰 신뢰도')
    assert hits[0]['location']=='문단 1'
    s.note(d['id'],'연구 목적','연구자의 해석',hits[:1],scope)
    s.note(d['id'],'연구 목적','다시 확인한 해석',hits[:1],scope,'연구자 확인')
    assert len(s.rows('document_notes'))==2
    with pytest.raises(ValueError):s.note(d['id'],'핵심 결과','허구',[],scope)
    with pytest.raises(ValueError):s.validate_evidence([dict(document_id=d['id'],location='문단 1',quote='없는 문장')],scope)
    with pytest.raises(ValueError):s.upload(s.files.get(d['sha256']),'가상.txt',d['source_id'])
    (s.files.root/d['sha256']).rename(s.files.root/'temporary-missing')
    with pytest.raises(FileNotFoundError):s.search(scope,'리뷰')


def test_backup_files_metadata_restore_and_corruption(tmp_path,library,source):
    s,d=seed(library,source);hits=s.search([d['id']],'리뷰');s.note(d['id'],'연구 질문','수동 해석',hits,[d['id']])
    blob=export_bundle(library);target=Library(tmp_path/'restored.db','actual')
    assert preview_bundle(target,blob)['added']['document_files']==1
    restore_bundle(target,blob,True)
    assert Documents(target).units(d['id'])==s.units(d['id'])
    assert Documents(target).rows('document_notes')==s.rows('document_notes')
    assert restore_bundle(target,blob,True)['document_files']==0
    with pytest.raises(ValueError):preview_bundle(Library(tmp_path/'example.db','example'),blob)
    out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as z,zipfile.ZipFile(out,'w') as corrupt:
        for name in z.namelist():corrupt.writestr(name,b'bad' if name==d['sha256'] else z.read(name))
    with pytest.raises(ValueError):restore_bundle(target,out.getvalue(),True)
    assert Documents(target).units(d['id'])==s.units(d['id'])


def test_ai_approval_exact_quotes_no_fabricated_success_and_budget(library,source,tmp_path):
    s,d=seed(library,source);hits=s.search([d['id']],'리뷰');payload,approval=prepare('어떤 근거인가?',hits)
    class Response:
        status_code=200
        def json(self):return {'choices':[{'message':{'content':json.dumps({'evidence':hits})}}]}
    class HTTP:
        calls=0
        def post(self,*a,**kw):self.calls+=1;return Response()
    http=HTTP()
    with pytest.raises(ValueError):analyze(s,payload,approval,'fake',False,http)
    with pytest.raises(ValueError):analyze(s,payload,'changed','fake',True,http)
    assert http.calls==0
    assert analyze(s,payload,approval,'fake',True,http)==hits
    hits[0]['quote']='날조된 근거'
    with pytest.raises(ValueError):analyze(s,payload,approval,'fake',True,http)
    for _ in range(10):reserve_budget(tmp_path,.01,1)
    with pytest.raises(ValueError):reserve_budget(tmp_path,.01,1)


def test_draft_diff_review_suggestions():
    changes=compare('서론\n기존 문장.\n결론','서론\n수정 문장.\n추가 문장.\n결론')
    assert changes and '수정' in [c['change'] for c in changes]
    assert any('인용 누락' in r['suggestion'] for r in suggestions('연구에 따르면 모든 사람이 90% 만족한다.'))


def test_v3_migration_failure_preserves_data_and_success_backup(library,source,monkeypatch):
    import sqlite3
    from reviewlab.services.research import Research
    Research(library).save(source)
    with library.connect() as db:
        db.execute('DROP TABLE document_notes');db.execute('DROP TABLE document_files')
        db.execute('DELETE FROM schema_migrations WHERE version=4')
    original=library.list_sources()
    read=Path.read_text
    def broken(path,*a,**kw):
        text=read(path,*a,**kw)
        return text+'; INVALID SQL;' if path.name=='004_documents.sql' else text
    with monkeypatch.context() as patch:
        patch.setattr(Path,'read_text',broken)
        with pytest.raises(sqlite3.Error):Library(library.path,'actual')
    with library.connect() as db:
        assert db.execute('SELECT max(version) FROM schema_migrations').fetchone()[0]==3
        assert not db.execute("SELECT name FROM sqlite_master WHERE name='document_files'").fetchone()
    assert library.list_sources()==original
    upgraded=Library(library.path,'actual')
    assert upgraded.list_sources()==original and not Documents(upgraded).rows()
    backups=list((library.path.parent/'backups').glob('*.sqlite3'))
    assert backups


def test_scan_blocks_analysis_claim_links_and_deleted_source_preserved(library,source):
    s,d=seed(library,source)
    scan=s.upload(pdf(),'scan.pdf',d['source_id'])
    with pytest.raises(ValueError,match='정상 추출'):s.search([scan['id']],'review')
    draft=s.upload('가상 주장: 리뷰 신뢰도'.encode(),'초고.txt',kind='draft')
    evidence=s.search([d['id'],draft['id']],'리뷰')
    n=s.note(draft['id'],'주장 검토','검토 제안',evidence,[d['id'],draft['id']])
    assert n['state']=='미확인'
    with library.connect() as db:db.execute('DELETE FROM sources WHERE id=?',(d['source_id'],))
    old=next(r for r in s.rows() if r['id']==d['id'])
    assert old['source_id'] is None and json.loads(old['reference_json'])['title']==source['title']
    assert s.units(d['id']) and s.rows('document_notes')


def test_ai_rejects_quote_outside_sent_scope_and_fabrication(library,source):
    s,d=seed(library,source);hits=s.search([d['id']],'리뷰');payload,approval=prepare('근거?',hits)
    class HTTP:
        def post(self,*a,**kw):
            class Response:
                status_code=200
                def json(self):return {'choices':[{'message':{'content':json.dumps({'evidence':[evidence]})}}]}
            return Response()
    evidence=dict(document_id=d['id'],location='문단 2',quote='표본 10명')
    with pytest.raises(ValueError,match='전송 범위 밖'):analyze(s,payload,approval,'fake',True,HTTP())
    evidence['quote']='1000명이 만족했다'
    with pytest.raises(ValueError,match='원문에서 확인되지 않음'):analyze(s,payload,approval,'fake',True,HTTP())


def test_database_only_cloud_migration_rejects_missing_file_transfer(library,source,tmp_path):
    from reviewlab.cloud.migration import prepare as migrate
    from reviewlab.cloud.library import CloudLibrary
    from reviewlab.cloud.client import CloudError
    from cloud_helpers import MemoryServer
    s,d=seed(library,source);server=MemoryServer()
    cloud=CloudLibrary(server.client(),server.workspace,'actual',tmp_path/'cloud')
    try:
        with pytest.raises(CloudError,match='파일 포함'):migrate(cloud,library.path)
        assert not cloud.dataset()['document_files']
    finally:cloud.close()
    assert s.units(d['id'])


def test_reject_forged_extracted_text_even_with_matching_hash(library,source):
    service,doc=seed(library,source);blob=export_bundle(library)
    with zipfile.ZipFile(io.BytesIO(blob)) as z:entries={name:z.read(name) for name in z.namelist()}
    manifest=json.loads(entries['manifest.json']);row=manifest['data']['document_files'][0]
    forged=packed(dict(extractor='reviewlab-v1',units=[dict(location='문단 1',text='원본에는 없는 허위 결과')]))
    del entries[row['text_sha256']];row['text_sha256']=digest(forged);entries[row['text_sha256']]=forged
    entries['manifest.json']=packed(manifest);out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        for name,value in entries.items():z.writestr(name,value)
    with pytest.raises(ValueError,match='원본과 추출'):preview_bundle(library,out.getvalue())
    assert service.units(doc['id'])[0]['text'].startswith('가상 연구 목적')


def test_korean_pdf_page_evidence(library,source):
    import matplotlib
    matplotlib.use('Agg')
    from reviewlab.analysis.charts import plt  # 기존 Python 3.12 호환 폰트 등록 재사용
    from matplotlib.backends.backend_pdf import PdfPages
    out=io.BytesIO()
    with matplotlib.rc_context({'pdf.fonttype':42}),PdfPages(out) as pages:
        for text in ('가상 연구 목적: 리뷰 신뢰도','가상 조사 대상: 대학생'):
            fig=plt.figure();fig.text(.1,.5,text);pages.savefig(fig);plt.close(fig)
    service,original=seed(library,source)
    document=service.upload(out.getvalue(),'가상.pdf',original['source_id'])
    assert document['status']=='ready'
    evidence=service.search([document['id']],'대학생')
    assert evidence[0]['location']=='페이지 2' and '대학생' in evidence[0]['quote']
    service.validate_evidence(evidence,[document['id']])


def test_complete_bundle_preserves_all_stages_and_histories(library,source,event,tmp_path):
    from test_surveys import seed as survey_seed
    from test_reviews import upload
    from reviewlab.services.research import Research
    from reviewlab.services.reviews import ReviewService
    from reviewlab.services.review_transfer import save_run
    from reviewlab.analysis.reviews import statistics,context,patterns
    from reviewlab.cloud.dataset import read_local
    service,doc=seed(library,source)
    Research(library).verify(doc['source_id'],event,True)
    Research(library).add_tag('가상 근거')
    tag=next(t['id'] for t in library.tags() if t['name']=='가상 근거')
    Research(library).save(source,[tag],doc['source_id'])
    reviews=ReviewService(library);pid=reviews.create_project('가상 복원 검증')
    upload(reviews,pid);rows=reviews.repo.rows(pid)
    reviews.note(rows[0]['id'],'가상 의견')
    reviews.check(rows[0]['id'],'검토 완료','가상 담당','2026-10-27','가상 근거')
    reviews.repo.link_sources(pid,[doc['source_id']])
    tables,_=statistics(rows);ctx=context(reviews.repo.project(pid),rows,rows,{})
    ctx['input_ids']=[r['id'] for r in rows]
    save_run(reviews.repo,pid,tables,ctx,patterns(rows))
    surveys,spid,_,_=survey_seed(library)
    surveys.save_analysis(spid,{'input_ids':surveys.frame(spid)[0]['id'].tolist()},{'가상 응답 수':4})
    service.note(doc['id'],'한계','가상 조건',service.search([doc['id']],'리뷰'),[doc['id']])
    before=read_local(library.path,'actual');assert all(before.values())
    target=Library(tmp_path/'all-stages/reviewlab.sqlite3','actual')
    blob=export_bundle(library);restore_bundle(target,blob,True)
    after=read_local(target.path,'actual')
    for table,rows in before.items():
        if table not in ('tags','source_tags'):assert after[table]==rows,table
    names=lambda data:{r['id']:r['name'] for r in data['tags']}
    linked=lambda data:{(r['source_id'],names(data)[r['tag_id']]) for r in data['source_tags']}
    assert linked(after)==linked(before)
    assert read_local(library.path,'actual')==before
    assert Documents(target).units(doc['id'])==service.units(doc['id'])
